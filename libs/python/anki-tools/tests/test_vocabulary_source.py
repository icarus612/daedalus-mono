"""Contract tests for ``anki_tools.vocabulary_source``.

Written from the module's public-surface and behavior contract alone. The
implementation under test is never read by this file's author.
"""

import csv
import sys
from pathlib import Path

import openpyxl
import pytest

from anki_tools.vocabulary_source import (
    EXPECTED_HEADERS,
    SHEETS,
    WorkbookError,
    build_parser,
    main,
    read_sheet,
)

REAL_WORKBOOK = Path("~/Downloads/russian_vocabulary.xlsx").expanduser()

EXPECTED_ROW_COUNTS = {
    "Nouns": 541,
    "Verbs": 179,
    "Adjectives": 152,
    "Adverbs": 152,
}


def make_workbook(path: Path, sheet_name: str, rows: list[list]) -> None:
    workbook = openpyxl.Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    sheet = workbook.create_sheet(sheet_name)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def adverbs_header() -> list[str]:
    return list(EXPECTED_HEADERS["Adverbs"])


class TestExpectedHeaders:
    def test_nouns(self):
        assert EXPECTED_HEADERS["Nouns"] == (
            "#",
            "Russian",
            "English",
            "Stress",
            "Gender",
            "Plural",
            "Genitive Sg",
            "Genitive Pl",
            "Animate",
            "Category",
            "Additional Info",
            "Status",
        )

    def test_verbs(self):
        assert EXPECTED_HEADERS["Verbs"] == (
            "#",
            "Imperfective",
            "Perfective",
            "English",
            "Stress",
            "Conjugation",
            "я (1sg)",
            "ты (2sg)",
            "они (3pl)",
            "Past (m / f)",
            "Case / Preposition",
            "Additional Info",
            "Status",
        )

    def test_adjectives(self):
        assert EXPECTED_HEADERS["Adjectives"] == (
            "#",
            "Russian (m)",
            "English",
            "Stress",
            "Feminine",
            "Neuter",
            "Plural",
            "Short Form",
            "Comparative",
            "Opposite",
            "Stem Type",
            "Additional Info",
            "Status",
        )

    def test_adverbs(self):
        assert EXPECTED_HEADERS["Adverbs"] == (
            "#",
            "Adverb",
            "English",
            "Stress",
            "From Adjective",
            "Formation",
            "Comparative",
            "Additional Info",
            "Status",
        )

    def test_sheets_tuple(self):
        assert SHEETS == ("Nouns", "Verbs", "Adjectives", "Adverbs")


@pytest.mark.skipif(
    not REAL_WORKBOOK.exists(), reason="real workbook not present locally"
)
class TestRealWorkbook:
    @pytest.mark.parametrize("sheet", SHEETS)
    def test_row_counts(self, sheet):
        rows = read_sheet(str(REAL_WORKBOOK), sheet)
        assert len(rows) == EXPECTED_ROW_COUNTS[sheet]


class TestHeaderValidation:
    def test_renamed_column_raises(self, tmp_path):
        header = adverbs_header()
        header[1] = "Adverb (renamed)"
        path = tmp_path / "bad_header.xlsx"
        make_workbook(
            path,
            "Adverbs",
            [header, ["1", "быстро", "quickly", "", "", "", "", "", ""]],
        )

        with pytest.raises(WorkbookError) as exc_info:
            read_sheet(str(path), "Adverbs")

        message = str(exc_info.value)
        assert "Adverbs" in message
        assert str(tuple(header)) in message
        assert str(EXPECTED_HEADERS["Adverbs"]) in message


class TestTabRejection:
    def test_tab_in_cell_raises(self, tmp_path):
        header = adverbs_header()
        row1 = ["1", "быстро", "quickly", "", "", "", "", "", ""]
        row2 = ["2", "медленно", "slow\tly", "", "", "", "", "", ""]
        path = tmp_path / "tab.xlsx"
        make_workbook(path, "Adverbs", [header, row1, row2])

        with pytest.raises(WorkbookError) as exc_info:
            read_sheet(str(path), "Adverbs")

        message = str(exc_info.value)
        assert "Adverbs" in message
        assert "2" in message
        assert "English" in message

    def test_cr_in_cell_raises(self, tmp_path):
        header = adverbs_header()
        row = ["1", "быстро", "quick\rly", "", "", "", "", "", ""]
        path = tmp_path / "cr.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        with pytest.raises(WorkbookError):
            read_sheet(str(path), "Adverbs")

    def test_lf_in_cell_raises(self, tmp_path):
        header = adverbs_header()
        row = ["1", "быстро", "quick\nly", "", "", "", "", "", ""]
        path = tmp_path / "lf.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        with pytest.raises(WorkbookError):
            read_sheet(str(path), "Adverbs")


class TestCellNormalization:
    def test_dash_only_cell_survives(self, tmp_path):
        header = adverbs_header()
        row = ["1", "быстро", "quickly", "—", "", "", "", "", ""]
        path = tmp_path / "dash.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        rows = read_sheet(str(path), "Adverbs")

        assert rows[0]["Stress"] == "—"

    @pytest.mark.parametrize("dash", ["—", "–", "-", "--"])
    def test_various_dash_forms_survive(self, tmp_path, dash):
        header = adverbs_header()
        row = ["1", "быстро", "quickly", dash, "", "", "", "", ""]
        path = tmp_path / "dashvariant.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        rows = read_sheet(str(path), "Adverbs")

        assert rows[0]["Stress"] == dash

    def test_whole_number_float_becomes_int_text(self, tmp_path):
        header = adverbs_header()
        row = [104.0, "быстро", "quickly", "", "", "", "", "", ""]
        path = tmp_path / "float.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        rows = read_sheet(str(path), "Adverbs")

        assert rows[0]["#"] == "104"

    def test_none_becomes_empty_string(self, tmp_path):
        header = adverbs_header()
        row = ["1", "быстро", "quickly", None, "", "", "", "", ""]
        path = tmp_path / "none.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        rows = read_sheet(str(path), "Adverbs")

        assert rows[0]["Stress"] == ""

    def test_string_is_stripped(self, tmp_path):
        header = adverbs_header()
        row = ["1", "  быстро  ", "quickly", "", "", "", "", "", ""]
        path = tmp_path / "strip.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        rows = read_sheet(str(path), "Adverbs")

        assert rows[0]["Adverb"] == "быстро"


class TestEmptyRowDropped:
    def test_all_empty_trailing_row_dropped(self, tmp_path):
        header = adverbs_header()
        row1 = ["1", "быстро", "quickly", "", "", "", "", "", ""]
        blank_row = ["", "", "", "", "", "", "", "", ""]
        path = tmp_path / "trailing_blank.xlsx"
        make_workbook(path, "Adverbs", [header, row1, blank_row])

        rows = read_sheet(str(path), "Adverbs")

        assert len(rows) == 1
        assert rows[0]["Adverb"] == "быстро"

    def test_returned_dicts_keyed_by_all_headers(self, tmp_path):
        header = adverbs_header()
        row = ["1", "быстро", "quickly", "", "от быстрый", "суффикс -о", "", "", ""]
        path = tmp_path / "keys.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        rows = read_sheet(str(path), "Adverbs")

        assert set(rows[0].keys()) == set(EXPECTED_HEADERS["Adverbs"])


class TestRoundTrip:
    def test_read_write_reread_matches(self, tmp_path):
        header = adverbs_header()
        row1 = [
            "1",
            "быстро",
            "quickly",
            "—",
            "от быстрый",
            "суффикс -о",
            "быстрее",
            "",
            "",
        ]
        row2 = ["2", "медленно", "slowly", "", "от медленный", "суффикс -о", "", "", ""]
        path = tmp_path / "roundtrip.xlsx"
        make_workbook(path, "Adverbs", [header, row1, row2])

        original_rows = read_sheet(str(path), "Adverbs")

        out_path = tmp_path / "adverbs.tsv"
        from anki_tools.vocabulary_source import write_tsv

        write_tsv(original_rows, EXPECTED_HEADERS["Adverbs"], str(out_path))

        with open(out_path, encoding="utf-8", newline="") as f:
            reread_rows = list(csv.DictReader(f, delimiter="\t"))

        assert reread_rows == original_rows

    def test_tsv_format(self, tmp_path):
        header = adverbs_header()
        row = ["1", "быстро", "quickly", "—", "", "", "", "", ""]
        path = tmp_path / "format.xlsx"
        make_workbook(path, "Adverbs", [header, row])

        rows = read_sheet(str(path), "Adverbs")

        out_path = tmp_path / "adverbs.tsv"
        from anki_tools.vocabulary_source import write_tsv

        write_tsv(rows, EXPECTED_HEADERS["Adverbs"], str(out_path))

        raw = out_path.read_bytes()
        assert b"\r\n" not in raw
        assert not raw.startswith(b"\xef\xbb\xbf")

        text = raw.decode("utf-8")
        lines = text.split("\n")
        assert lines[0] == "\t".join(EXPECTED_HEADERS["Adverbs"])


class TestMainDryRun:
    def build_synthetic_workbook(self, tmp_path):
        path = tmp_path / "vocab.xlsx"
        workbook = openpyxl.Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        for sheet_name in SHEETS:
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(list(EXPECTED_HEADERS[sheet_name]))
            row = ["1"] + [""] * (len(EXPECTED_HEADERS[sheet_name]) - 1)
            sheet.append(row)
        workbook.save(path)
        return path

    def test_dry_run_prints_counts_and_writes_nothing(
        self, tmp_path, monkeypatch, capsys
    ):
        workbook_path = self.build_synthetic_workbook(tmp_path)
        out_dir = tmp_path / "out"

        monkeypatch.setattr(
            sys,
            "argv",
            [
                "vocabulary_source",
                "--workbook",
                str(workbook_path),
                "--out-dir",
                str(out_dir),
                "--dry-run",
            ],
        )

        with pytest.raises(SystemExit) as exc_info:
            main()

        assert exc_info.value.code == 0

        captured = capsys.readouterr()
        for sheet in SHEETS:
            assert f"{sheet}: 1 rows" in captured.out

        assert not out_dir.exists() or not any(out_dir.iterdir())

    def test_workbook_error_exits_one(self, tmp_path, monkeypatch, capsys):
        header = adverbs_header()
        header[1] = "Wrong Column"
        path = tmp_path / "bad.xlsx"
        make_workbook(path, "Adverbs", [header])

        workbook = openpyxl.load_workbook(path)
        for sheet_name in SHEETS:
            if sheet_name not in workbook.sheetnames:
                sheet = workbook.create_sheet(sheet_name)
                sheet.append(list(EXPECTED_HEADERS[sheet_name]))
        workbook.save(path)

        monkeypatch.setattr(
            sys,
            "argv",
            ["vocabulary_source", "--workbook", str(path), "--dry-run"],
        )

        with pytest.raises(SystemExit) as exc_info:
            main()

        assert exc_info.value.code == 1


def workbook_arg(args):
    for name in ("workbook", "workbook_path"):
        if hasattr(args, name):
            return getattr(args, name)
    raise AssertionError("parsed args expose no workbook-path attribute")


class TestBuildParser:
    def test_defaults(self):
        parser = build_parser()
        args = parser.parse_args([])

        assert workbook_arg(args) == str(
            Path("~/Downloads/russian_vocabulary.xlsx").expanduser()
        )
        assert not args.dry_run

    def test_out_dir_default_next_to_module(self):
        parser = build_parser()
        args = parser.parse_args([])

        assert Path(args.out_dir).name == "russian-vocabulary"
        assert Path(args.out_dir).parent.name == "data"

    def test_explicit_flags(self):
        parser = build_parser()
        args = parser.parse_args(
            [
                "--workbook",
                "/tmp/foo.xlsx",
                "--out-dir",
                "/tmp/out",
                "--dry-run",
            ]
        )

        assert workbook_arg(args) == "/tmp/foo.xlsx"
        assert args.out_dir == "/tmp/out"
        assert args.dry_run
